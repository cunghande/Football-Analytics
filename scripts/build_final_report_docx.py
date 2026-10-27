"""Build the Vietnamese academic report from its Markdown source.

The Markdown file is the editable source of truth; this script needs python-docx.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "report/BAO_CAO_HOAN_CHINH.md"
TARGET = ROOT / "report/BAO_CAO_HOAN_CHINH.docx"


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    tc_pr.append(borders)


def add_inline(paragraph, source):
    """Render the subset of inline Markdown used by the report."""
    pattern = re.compile(r"(\[([^]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`|\*([^*]+)\*)")
    cursor = 0
    for match in pattern.finditer(source):
        if match.start() > cursor:
            paragraph.add_run(source[cursor:match.start()])
        if match.group(2):
            run = paragraph.add_run(match.group(2))
            run.underline = True
        elif match.group(4):
            run = paragraph.add_run(match.group(4))
            run.bold = True
        elif match.group(5):
            run = paragraph.add_run(match.group(5))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            run = paragraph.add_run(match.group(6))
            run.italic = True
        cursor = match.end()
    if cursor < len(source):
        paragraph.add_run(source[cursor:])


def add_page_field(paragraph):
    paragraph.add_run("Trang ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build():
    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.16

    for name, size, before, after in (
        ("Title", 18, 0, 14),
        ("Heading 1", 14, 13, 7),
        ("Heading 2", 12, 10, 5),
    ):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        p_pr = style.element.get_or_add_pPr()
        for border in p_pr.findall(qn("w:pBdr")):
            p_pr.remove(border)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer)
    footer.style = normal

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    idx = 0
    while idx < len(lines):
        line = lines[idx].strip()
        if not line:
            idx += 1
            continue
        if line.startswith("```"):
            code = []
            idx += 1
            while idx < len(lines) and not lines[idx].startswith("```"):
                code.append(lines[idx])
                idx += 1
            para = document.add_paragraph()
            para.paragraph_format.space_after = Pt(9)
            for n, code_line in enumerate(code):
                run = para.add_run(code_line + ("\n" if n < len(code) - 1 else ""))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
            idx += 1
            continue
        if line.startswith("| "):
            rows = []
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append(cells)
                idx += 1
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.autofit = True
            for ri, row in enumerate(rows):
                for ci, value in enumerate(row):
                    cell = table.cell(ri, ci)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    add_inline(p, value)
                    for run in p.runs:
                        run.font.size = Pt(8.5 if len(row) >= 5 else 9)
                        if ri == 0:
                            run.bold = True
                    set_cell_border(cell)
            continue
        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image_match:
            path = SOURCE.parent / image_match.group(2)
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(1)
            p.add_run().add_picture(str(path), width=Cm(15.8))
            idx += 1
            continue
        if line.startswith("# "):
            document.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            document.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 2")
        elif re.match(r"^\d+\. ", line):
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            add_inline(p, line)
        elif line.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        else:
            p = document.add_paragraph()
            if line.startswith("*Hình") or line.startswith("*Bảng"):
                p.paragraph_format.keep_with_next = False
                p.paragraph_format.space_after = Pt(9)
            add_inline(p, line)
        idx += 1

    document.core_properties.title = "Phát triển hệ thống phân tích bóng đá từ video bằng thị giác máy tính"
    document.core_properties.subject = "Báo cáo bài tập lớn thị giác máy tính"
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    build()
